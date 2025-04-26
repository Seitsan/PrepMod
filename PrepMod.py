import pandas as pd
import re
import os
import numpy as np
from tkinter import filedialog, Tk, messagebox
from rdkit import Chem
from rdkit.Chem import Descriptors, MolToMolBlock, SDWriter, rdMolDescriptors
from rdkit import RDLogger
from openpyxl import load_workbook
from sklearn.model_selection import KFold
import shutil

# Отключаем предупреждения RDKit
RDLogger.DisableLog('rdApp.*')

def clean_filename(text):
    """Создает валидное имя файла из строки"""
    text = re.sub(r'[\\/*?:"<>|]', "_", str(text))
    text = re.sub(r'\s+', '_', text)
    return text.strip('_')

def select_input_file():
    """Выбор входного файла через диалоговое окно"""
    root = Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    file_path = filedialog.askopenfilename(
        title="Выберите файл с данными",
        filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")]
    )
    root.destroy()
    return file_path

def find_cas_smiles_mapping(file_path):
    """Поиск соответствий CAS-SMILES во всех листах Excel"""
    cas_smiles_map = {}
    try:
        wb = load_workbook(file_path, read_only=True)
        
        for sheet_name in wb.sheetnames:
            try:
                df = pd.read_excel(file_path, sheet_name=sheet_name, dtype=str)
                
                cols_lower = [col.lower() for col in df.columns]
                
                if 'cas' in cols_lower and 'canonical smiles' in cols_lower:
                    cas_col = df.columns[cols_lower.index('cas')]
                    smiles_col = df.columns[cols_lower.index('canonical smiles')]
                    
                    for _, row in df.iterrows():
                        cas = str(row[cas_col]).strip() if pd.notna(row[cas_col]) else None
                        smiles = str(row[smiles_col]).strip() if pd.notna(row[smiles_col]) else None
                        
                        if cas and smiles and cas not in cas_smiles_map:
                            mol = Chem.MolFromSmiles(smiles)
                            if mol is not None:
                                cas_smiles_map[cas] = smiles
            except Exception as e:
                print(f"Ошибка при обработке листа {sheet_name}: {str(e)}")
                continue
    
    except Exception as e:
        print(f"Ошибка при чтении файла: {str(e)}")
        return {}
    
    return cas_smiles_map

def load_and_process_data(file_path):
    """Загрузка и обработка данных с поиском SMILES"""
    cas_smiles_map = find_cas_smiles_mapping(file_path)
    
    if not cas_smiles_map:
        print("Не найдено соответствий CAS-SMILES в файле")
        return None
    
    try:
        df = pd.read_excel(file_path, sheet_name=0)
    except Exception as e:
        print(f"Ошибка при загрузке главного листа: {str(e)}")
        return None
    
    required_columns = ['CAS', 'Effect value', 'Effect', 'Test statistic', 'Duration (hours)', 'Latin name']
    missing_cols = [col for col in required_columns if col not in df.columns]
    
    if missing_cols:
        print(f"Отсутствуют обязательные столбцы: {', '.join(missing_cols)}")
        return None
    
    df['CAS'] = df['CAS'].astype(str).str.strip()
    df['Structure'] = np.nan
    df['Canonical SMILES'] = np.nan
    
    found_count = 0
    for idx, row in df.iterrows():
        cas = row['CAS']
        if cas in cas_smiles_map:
            df.at[idx, 'Structure'] = cas_smiles_map[cas]
            df.at[idx, 'Canonical SMILES'] = cas_smiles_map[cas]
            found_count += 1
    
    print(f"\nНайдено SMILES для {found_count} из {len(df)} записей")
    
    initial_count = len(df)
    df = df.dropna(subset=['Structure'])
    removed_count = initial_count - len(df)
    
    if removed_count > 0:
        print(f"Удалено записей без SMILES: {removed_count}")
    
    if df.empty:
        print("Нет данных для обработки после фильтрации")
        return None
    
    if 'Records_combined' not in df.columns:
        df['Records_combined'] = 1
    
    return df

def save_to_excel(dataframe, default_name):
    """Сохранение DataFrame в Excel"""
    root = Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    save_path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
        title="Сохранить файл как",
        initialfile=default_name
    )
    root.destroy()

    if save_path:
        try:
            dataframe.to_excel(save_path, index=False)
            print(f"Файл успешно сохранён: {save_path}")
            return True
        except Exception as e:
            print(f"Ошибка при сохранении файла: {e}")
            return False
    print("Сохранение отменено")
    return False

def show_top_combinations(combination_counts):
    """Выводит топ-10 сочетаний параметров"""
    print("\nТоп-10 сочетаний Effect-Test-Duration:")
    print("-" * 85)
    print(f"{'Effect':<40} | {'Test Statistic':<20} | {'Duration (h)':>12} | {'Count':>6}")
    print("-" * 85)

    for _, row in combination_counts.head(10).iterrows():
        effect = (row['Effect'][:37] + '...') if len(row['Effect']) > 40 else row['Effect']
        test = (row['Test statistic'][:17] + '...') if len(row['Test statistic']) > 20 else row['Test statistic']
        print(f"{effect:<40} | {test:<20} | {row['Duration (hours)']:>12} | {row['Count']:>6}")
    print("-" * 85)

def filter_by_count(df, combination_counts):
    """Фильтрация по количеству сочетаний"""
    try:
        count = int(input("Введите минимальное количество записей (>=): "))
        filtered = combination_counts[combination_counts['Count'] >= count]

        if not filtered.empty:
            mask = df.set_index(['Effect', 'Test statistic', 'Duration (hours)']).index.isin(
                filtered.set_index(['Effect', 'Test statistic', 'Duration (hours)']).index)
            return df[mask]
        print(f"Нет записей с количеством >= {count}")
    except ValueError:
        print("Пожалуйста, введите целое число")
    return pd.DataFrame()

def filter_by_effect(df, effect_counts):
    """Фильтрация только по Effect"""
    print("\nДоступные эффекты:")
    for i, (effect, count) in enumerate(effect_counts.items(), 1):
        print(f"{i}. {effect} ({count} записей)")

    effect_input = input("\nВведите номер или название Effect: ")
    try:
        effect_num = int(effect_input)
        if 1 <= effect_num <= len(effect_counts):
            effect = effect_counts.index[effect_num - 1]
            return df[df['Effect'] == effect]
    except ValueError:
        if effect_input in effect_counts.index:
            return df[df['Effect'] == effect_input]

    print("Указанный эффект не найден")
    return pd.DataFrame()

def filter_by_effect_test(df):
    """Фильтрация по сочетанию Effect-Test statistic"""
    combinations = df.groupby(['Effect', 'Test statistic']).size().reset_index(name='Count')
    
    print("\nДоступные комбинации Effect - Test statistic:")
    for i, row in combinations.iterrows():
        print(f"{i+1}. {row['Effect']} - {row['Test statistic']} ({row['Count']} записей)")
    
    selected = input("\nВведите номера комбинаций через запятую (или 'all' для всех): ")
    
    if selected.lower() == 'all':
        return df
    else:
        try:
            selected_indices = [int(x.strip()) - 1 for x in selected.split(',')]
            selected_combinations = combinations.iloc[selected_indices]
            
            mask = df.apply(lambda x: any(
                (x['Effect'] == combo['Effect']) and 
                (x['Test statistic'] == combo['Test statistic'])
                for _, combo in selected_combinations.iterrows()
            ), axis=1)
            
            return df[mask]
        except Exception as e:
            print(f"Ошибка при выборе комбинаций: {e}")
            return pd.DataFrame()

def filter_by_combination(df):
    """Фильтрация по полному сочетанию параметров"""
    print("\nПример ввода: Growth inhibition, t-test, 24")
    effect = input("Введите Effect: ")
    test = input("Введите Test statistic: ")
    try:
        duration = int(input("Введите Duration (hours): "))
        mask = (df['Effect'] == effect) & \
               (df['Test statistic'] == test) & \
               (df['Duration (hours)'] == duration)
        return df[mask]
    except ValueError:
        print("Пожалуйста, введите число для Duration")
    return pd.DataFrame()

def apply_count_filter(df, group_columns):
    """Применяет фильтр по количеству записей"""
    combo_counts = df.groupby(group_columns).size().reset_index(name='Count')
    
    print("\nВыберите действие:")
    print("1. Работать со всеми выбранными комбинациями")
    print("2. Оставить только комбинации с количеством записей >= N")
    
    while True:
        choice = input("Ваш выбор: ").strip()
        
        if choice == '1':
            return df
        elif choice == '2':
            try:
                min_count = int(input("Введите минимальное количество записей: "))
                filtered_combos = combo_counts[combo_counts['Count'] >= min_count]
                
                if filtered_combos.empty:
                    print("Нет комбинаций с указанным количеством записей")
                    return pd.DataFrame()
                
                mask = df.set_index(group_columns).index.isin(
                    filtered_combos.set_index(group_columns).index)
                return df[mask]
            except ValueError:
                print("Пожалуйста, введите целое число")
        else:
            print("Неверный выбор")

def prepare_for_modeling(dataframe):
    """Подготовка данных для моделирования с кросс-валидацией"""
    print("\n=== ПОДГОТОВКА ДАННЫХ ДЛЯ МОДЕЛИРОВАНИЯ ===")

    required_columns = ['CAS', 'Effect value', 'Structure', 'Effect', 'Test statistic', 'Duration (hours)', 'Latin name']
    if not all(col in dataframe.columns for col in required_columns):
        missing = [col for col in required_columns if col not in dataframe.columns]
        print(f"ОШИБКА: Отсутствуют столбцы {', '.join(missing)}")
        return pd.DataFrame()

    original_count = len(dataframe)
    print(f"\nИсходное количество записей: {original_count}")

    # Удаление дубликатов
    grouped = dataframe.groupby(['Effect', 'Test statistic', 'Duration (hours)', 'CAS'])
    aggregated_df = grouped.agg({
        'Chemical name': 'first',
        'Structure': 'first',
        'Effect value': 'median',
        'Latin name': 'first',
        'Records_combined': 'sum' if 'Records_combined' in dataframe.columns else 'size'
    }).reset_index()

    duplicates_removed = original_count - len(aggregated_df)
    print(f"\n[ДУБЛИКАТЫ] Удалено записей: {duplicates_removed}")

    # Конвертация SMILES в Molfile
    print("\nКонвертация SMILES в Molfile и расчет свойств...")
    valid_mols = []
    mol_props = []

    for _, row in aggregated_df.iterrows():
        try:
            mol = Chem.MolFromSmiles(row['Structure'])
            if mol is not None:
                mol_weight = Descriptors.MolWt(mol)
                formal_charge = Chem.GetFormalCharge(mol)
                num_components = len(Chem.GetMolFrags(mol))
                num_carbons = len([atom for atom in mol.GetAtoms() if atom.GetAtomicNum() == 6])
                
                props = {
                    'Molecular_weight': round(mol_weight, 3),
                    'Formal_charge': formal_charge,
                    'Num_components': num_components,
                    'Num_carbons': num_carbons,
                    'Original_index': _,
                    'Latin_name': row['Latin name'],
                    'Effect': row['Effect'],
                    'Test_statistic': row['Test statistic'],
                    'Duration_hours': row['Duration (hours)'],
                    'Effect_value': row['Effect value']
                }
                
                mol_props.append(props)
                valid_mols.append(mol)
        except Exception as e:
            print(f"Ошибка обработки структуры {row['Structure']}: {e}")

    props_df = pd.DataFrame(mol_props)
    print(f"\nУспешно обработано {len(props_df)} структур")

    # Применение фильтров
    print("\nПрименение фильтров:")
    
    before = len(props_df)
    props_df = props_df[props_df['Num_components'] == 1]
    print(f"- Удалено не монокомпонентных структур: {before - len(props_df)}")
    
    before = len(props_df)
    props_df = props_df[props_df['Formal_charge'] == 0]
    print(f"- Удалено не электронейтральных структур: {before - len(props_df)}")
    
    before = len(props_df)
    props_df = props_df[props_df['Molecular_weight'] < 1250]
    print(f"- Удалено структур с MW > 1250: {before - len(props_df)}")
    
    before = len(props_df)
    props_df = props_df[props_df['Num_carbons'] >= 3]
    print(f"- Удалено структур с <3 атомами C: {before - len(props_df)}")

    # Создание финальных структур
    final_mols = [valid_mols[i] for i in props_df['Original_index']]
    for mol, props in zip(final_mols, props_df.to_dict('records')):
        for key, value in props.items():
            if key != 'Original_index':
                mol.SetProp(key, str(value))

    apply_count_filter(props_df, ['Latin_name', 'Effect', 'Test_statistic', 'Duration_hours'])

    # Подготовка к кросс-валидации
    root = Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    output_dir = filedialog.askdirectory(title="Выберите папку для сохранения результатов")
    root.destroy()

    if not output_dir:
        print("Сохранение отменено")
        return pd.DataFrame()

    print("\nПодготовка к кросс-валидации...")
    
    grouped_data = props_df.groupby(['Latin_name', 'Effect', 'Test_statistic', 'Duration_hours'])
    kf = KFold(n_splits=5, shuffle=True, random_state=42)
    
    # Структуры для хранения данных
    fold_train_data = [[] for _ in range(5)]
    fold_test_data = [[] for _ in range(5)]
    general_groups = {}

    for _, group in grouped_data:
        indices = group.index.tolist()
        key = (group.iloc[0]['Latin_name'], group.iloc[0]['Effect'],
               group.iloc[0]['Test_statistic'], group.iloc[0]['Duration_hours'])
        
        general_groups[key] = [valid_mols[i] for i in indices]
        
        if len(indices) >= 5:
            for fold, (train_idx, test_idx) in enumerate(kf.split(indices)):
                for idx in [indices[i] for i in train_idx]:
                    fold_train_data[fold].append(valid_mols[idx])
                for idx in [indices[i] for i in test_idx]:
                    fold_test_data[fold].append(valid_mols[idx])
        else:
            print(f"Пропущена группа {key} - недостаточно данных")

    # Сохранение данных
    print("\nСохранение результатов...")
    
    # Сохраняем общие SDF
    for (latin_name, effect, test_stat, duration), mols in general_groups.items():
        current_dir = os.path.join(
            output_dir,
            clean_filename(latin_name),
            clean_filename(effect),
            clean_filename(test_stat),
            f"{duration}h"
        )
        os.makedirs(current_dir, exist_ok=True)
        
        sdf_path = os.path.join(current_dir, "data.sdf")
        writer = SDWriter(sdf_path)
        for mol in mols:
            writer.write(mol)
        writer.close()
        print(f"Сохранен общий файл: {sdf_path}")

    # Сохраняем train/test для каждого фолда
    for fold_idx in range(5):
        fold_train_groups = {}
        fold_test_groups = {}
        
        for mol in fold_train_data[fold_idx]:
            key = (mol.GetProp('Latin_name'), mol.GetProp('Effect'),
                   mol.GetProp('Test_statistic'), mol.GetProp('Duration_hours'))
            if key not in fold_train_groups:
                fold_train_groups[key] = []
            fold_train_groups[key].append(mol)
            
        for mol in fold_test_data[fold_idx]:
            key = (mol.GetProp('Latin_name'), mol.GetProp('Effect'),
                   mol.GetProp('Test_statistic'), mol.GetProp('Duration_hours'))
            if key not in fold_test_groups:
                fold_test_groups[key] = []
            fold_test_groups[key].append(mol)
        
        for (latin_name, effect, test_stat, duration) in fold_train_groups.keys():
            fold_dir = os.path.join(
                output_dir,
                clean_filename(latin_name),
                clean_filename(effect),
                clean_filename(test_stat),
                f"{duration}h",
                f"Fold_{fold_idx+1}"
            )
            os.makedirs(fold_dir, exist_ok=True)
            
            # Сохраняем train
            train_path = os.path.join(fold_dir, "train.sdf")
            writer = SDWriter(train_path)
            for mol in fold_train_groups.get((latin_name, effect, test_stat, duration), []):
                writer.write(mol)
            writer.close()
            
            # Сохраняем test
            test_path = os.path.join(fold_dir, "test.sdf")
            writer = SDWriter(test_path)
            for mol in fold_test_groups.get((latin_name, effect, test_stat, duration), []):
                writer.write(mol)
            writer.close()
            
            print(f"Сохранен фолд {fold_idx+1}: {train_path}, {test_path}")

    print(f"\nРезультаты сохранены в: {output_dir}")
    print(f"Структура папок:")
    print(f"LatinName/Effect/TestStatistic/Duration/")
    print(f"  ├── data.sdf")
    print(f"  ├── Fold_1/")
    print(f"  │   ├── train.sdf")
    print(f"  │   └── test.sdf")
    print(f"  ├── Fold_2/")
    print(f"  │   ├── train.sdf")
    print(f"  │   └── test.sdf")
    print(f"  └── ...")
    print(f"Всего фолдов: 5")
    
    # Возвращаем DataFrame с отфильтрованными данными
    result_df = aggregated_df.loc[props_df['Original_index']].reset_index(drop=True)
    result_df['Structure'] = [MolToMolBlock(mol) for mol in final_mols]
    
    print(f"\n=== ИТОГИ ОБРАБОТКИ ===")
    print(f"Обработано записей: {len(result_df)} из {original_count}")
    print(f"Всего удалено: {original_count - len(result_df)} записей")

    return result_df

def save_split_by_combination(result_df, prefix=""):
    """Сохранение с разделением по сочетаниям параметров"""
    grouped = result_df.groupby(['Latin name', 'Effect', 'Test statistic', 'Duration (hours)'])
    saved_files = 0

    for (latin_name, effect, test, duration), group in grouped:
        filename = f"{prefix}{clean_filename(latin_name)}_{clean_filename(effect)}_{clean_filename(test)}_{duration}h.xlsx"
        if save_to_excel(group, filename):
            saved_files += 1

    print(f"\nСохранено файлов: {saved_files}")
    return saved_files > 0

def save_to_sdf_by_combination(dataframe):
    """Сохраняет подготовленные данные в SDF файлы с разделением по комбинациям"""
    if len(dataframe) == 0:
        print("Нет данных для сохранения")
        return False

    root = Tk()
    root.withdraw()
    root.attributes('-topmost', True)
    output_dir = filedialog.askdirectory(title="Выберите папку для сохранения SDF файлов")
    root.destroy()

    if not output_dir:
        print("Сохранение отменено")
        return False

    os.makedirs(output_dir, exist_ok=True)
    grouped = dataframe.groupby(['Latin name', 'Effect', 'Test statistic', 'Duration (hours)'])
    saved_files = 0

    for (latin_name, effect, test, duration), group in grouped:
        dir_path = os.path.join(
            output_dir,
            clean_filename(latin_name),
            clean_filename(effect),
            clean_filename(test),
            f"{duration}h"
        )
        os.makedirs(dir_path, exist_ok=True)
        
        filename = "data.sdf"
        filepath = os.path.join(dir_path, filename)

        try:
            writer = SDWriter(filepath)
            for _, row in group.iterrows():
                mol = Chem.MolFromMolBlock(row['Structure'])
                if mol is not None:
                    for col in group.columns:
                        if col != 'Structure' and pd.notna(row[col]):
                            mol.SetProp(col, str(row[col]))
                    writer.write(mol)
            writer.close()
            saved_files += 1
            print(f"Сохранен файл: {filepath}")
        except Exception as e:
            print(f"Ошибка при сохранении файла {filepath}: {e}")

    print(f"\nВсего сохранено SDF файлов: {saved_files}")
    return saved_files > 0

def analyze_and_filter_effects():
    """Основная функция программы"""
    try:
        from rdkit import Chem
        from rdkit.Chem import Descriptors
    except ImportError as e:
        print("\nОШИБКА: Необходимо установить RDKit")
        print("Установите: conda install -c conda-forge rdkit")
        return

    while True:
        print("\n" + "=" * 50)
        print("Анализ данных химических экспериментов")
        print("=" * 50)
        file_path = select_input_file()

        if not file_path:
            print("Файл не выбран.")
            break

        try:
            df = load_and_process_data(file_path)
            if df is None:
                continue

            while True:
                effect_counts = df['Effect'].value_counts().sort_values(ascending=False)
                combination_counts = df.groupby(['Effect', 'Test statistic', 'Duration (hours)']).size() \
                    .reset_index(name='Count') \
                    .sort_values('Count', ascending=False)

                print("\n" + "=" * 50)
                print("ГЛАВНОЕ МЕНЮ")
                print("1 - Показать статистику")
                print("2 - Фильтрация записей")
                print("3 - Подготовить данные для моделирования")
                print("4 - Выбрать другой файл")
                print("0 - Выход")
                main_choice = input("Ваш выбор: ").strip()

                if main_choice == '0':
                    return
                elif main_choice == '4':
                    break
                elif main_choice == '1':
                    print("\nКоличество записей по эффектам:")
                    for effect, count in effect_counts.items():
                        print(f"{effect}: {count} записей")
                    show_top_combinations(combination_counts)
                    
                    input("\nНажмите Enter чтобы продолжить...")
                elif main_choice == '2':
                    print("\nТип фильтра:")
                    print("1 - По количеству сочетаний Effect-Test-Duration (>=)")
                    print("2 - По Effect или сочетанию параметров")
                    filter_choice = input("Ваш выбор: ").strip()

                    result_df = pd.DataFrame()
                    if filter_choice == '1':
                        result_df = filter_by_count(df, combination_counts)
                    elif filter_choice == '2':
                        print("\nТип фильтра:")
                        print("1 - Только по Effect")
                        print("2 - По сочетанию Effect-Test statistic")
                        print("3 - По сочетанию Effect-Test-Duration")
                        sub_choice = input("Ваш выбор: ").strip()

                        if sub_choice == '1':
                            result_df = filter_by_effect(df, effect_counts)
                            if not result_df.empty:
                                result_df = apply_count_filter(result_df, ['Effect'])
                        elif sub_choice == '2':
                            result_df = filter_by_effect_test(df)
                            if not result_df.empty:
                                result_df = apply_count_filter(result_df, ['Effect', 'Test statistic'])
                        elif sub_choice == '3':
                            result_df = filter_by_combination(df)
                            if not result_df.empty:
                                result_df = apply_count_filter(result_df, ['Effect', 'Test statistic', 'Duration (hours)'])
                        else:
                            print("Неверный выбор")
                    else:
                        print("Неверный выбор")

                    if not result_df.empty:
                        while True:
                            print(f"\nНайдено {len(result_df)} записей")
                            print("1 - Сохранить в Excel")
                            print("2 - Подготовить для моделирования")
                            print("0 - Назад")
                            action_choice = input("Ваш выбор: ").strip()

                            if action_choice == '0':
                                break
                            elif action_choice == '1':
                                print("\nФормат сохранения:")
                                print("1 - Сохранить все данные в один файл")
                                print("2 - Разделить по сочетаниям Latin-Эффект-Тест-Длительность")
                                save_choice = input("Ваш выбор: ").strip()

                                if save_choice == '1':
                                    default_name = "filtered_data.xlsx"
                                    save_to_excel(result_df, default_name)
                                elif save_choice == '2':
                                    save_split_by_combination(result_df)
                            elif action_choice == '2':
                                prepared_data = prepare_for_modeling(result_df)

                                if not prepared_data.empty:
                                    print("\nХотите сохранить подготовленные данные?")
                                    print("1 - Да, разделить по Latin-Эффект-Тест-Длительность (Excel)")
                                    print("2 - Да, сохранить как единый файл (Excel)")
                                    print("3 - Да, сохранить в SDF (раздельные файлы)")
                                    print("4 - Нет, продолжить без сохранения")
                                    save_choice = input("Ваш выбор: ").strip()

                                    if save_choice == '1':
                                        save_split_by_combination(prepared_data, "prepared_")
                                    elif save_choice == '2':
                                        save_to_excel(prepared_data, "prepared_for_modeling.xlsx")
                                    elif save_choice == '3':
                                        save_to_sdf_by_combination(prepared_data)
                elif main_choice == '3':
                    prepared_data = prepare_for_modeling(df)
                    if not prepared_data.empty:
                        print("\nХотите сохранить подготовленные данные?")
                        print("1 - Да, разделить по Latin-Эффект-Тест-Длительность (Excel)")
                        print("2 - Да, сохранить как единый файл (Excel)")
                        print("3 - Да, сохранить в SDF (раздельные файлы)")
                        print("4 - Нет, продолжить без сохранения")
                        save_choice = input("Ваш выбор: ").strip()

                        if save_choice == '1':
                            save_split_by_combination(prepared_data, "prepared_")
                        elif save_choice == '2':
                            save_to_excel(prepared_data, "prepared_for_modeling.xlsx")
                        elif save_choice == '3':
                            save_to_sdf_by_combination(prepared_data)
                else:
                    print("Неверный выбор")
        except Exception as e:
            print(f"Ошибка при обработке файла: {e}")
            if not messagebox.askyesno("Ошибка", "Произошла ошибка. Продолжить работу?"):
                return

if __name__ == "__main__":
    analyze_and_filter_effects()